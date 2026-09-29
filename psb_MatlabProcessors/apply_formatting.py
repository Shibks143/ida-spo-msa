"""
apply_formatting.py
Called by compute_EDP_SideBySide_final.m to apply all cell formatting
to the output XLSX — colours, borders, fonts, fills — producing output
identical to EDP_SideBySide_v9.xlsx.

Usage (from MATLAB):
    pyrunfile('apply_formatting.py', out_file=out_file)

Or standalone:
    python apply_formatting.py "path\to\\EDP_SideBySide_MATLAB.xlsx"
"""

import sys
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Accept path from MATLAB pyrunfile or command line ─────────────────────
if 'out_file' not in dir():          # running standalone
    if len(sys.argv) < 2:
        raise SystemExit("Usage: python apply_formatting.py <xlsx_path>")
    out_file = sys.argv[1]

# ── Style helpers ─────────────────────────────────────────────────────────
def fill(h):   return PatternFill("solid", fgColor=h)
def font(bold=False, size=9, color="000000", name="Arial"):
    return Font(bold=bold, size=size, color=color, name=name)
def align(h="center", v="center", wrap=False):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

THICK = Side(style="thick", color="1F3864")
THIN  = Side(style="thin",  color="999999")
NONE  = Side(style=None)

def bd_orig():  return Border(left=THICK, right=THIN,  top=THIN, bottom=THIN)
def bd_curt():  return Border(left=THIN,  right=THIN,  top=THIN, bottom=THIN)
def bd_delta(): return Border(left=THIN,  right=THICK, top=THIN, bottom=THIN)

# Colours
F_TITLE    = fill("1F3864")
F_EQ_HDR   = fill("2E4057")
F_EDP_SEC  = fill("2C5F8A")
F_ORIG_HDR = fill("D6DCE4")
F_CURT_HDR = fill("D6DCE4")
F_DELT_HDR = fill("BDD0E8")
F_STORY    = fill("EEF2F7")
F_WHITE    = fill("FFFFFF")
F_POS      = fill("C6EFCE")
F_NEG      = fill("FFCCCC")
F_NEARZERO = fill("FFF2CC")

THRESH = 0.05
def delta_fill(d):
    if d is None:   return F_NEARZERO
    if d >  THRESH: return F_POS
    if d < -THRESH: return F_NEG
    return F_NEARZERO

FT_TITLE   = font(bold=True,  size=11, color="FFFFFF")
FT_EQ_HDR  = font(bold=True,  size=9,  color="FFFFFF")
FT_SUB_HDR = font(bold=True,  size=8,  color="1F3864")
FT_EDP_SEC = font(bold=True,  size=10, color="FFFFFF")
FT_STORY   = font(bold=True,  size=9,  color="2C3E50")
FT_DATA    = font(bold=False, size=9,  color="000000")

N_EQ        = 60
FIXED       = 2
COLS_PER_EQ = 3
TOTAL_COLS  = FIXED + N_EQ * COLS_PER_EQ

def col_o(j): return FIXED + 1 + j * COLS_PER_EQ
def col_c(j): return FIXED + 2 + j * COLS_PER_EQ
def col_d(j): return FIXED + 3 + j * COLS_PER_EQ

def wc(ws, r, c, fl=None, ft=None, al=None, fmt=None, bd=None):
    cell = ws.cell(row=r, column=c)
    if fl:  cell.fill          = fl
    if ft:  cell.font          = ft
    if al:  cell.alignment     = al
    if fmt: cell.number_format = fmt
    if bd:  cell.border        = bd

def apply_eq_merged_border(ws, row, j):
    """Clean single-box border on merged EQ cell — no internal lines."""
    o, c, d = col_o(j), col_c(j), col_d(j)
    ws.cell(row=row, column=o).border = Border(left=THICK, right=NONE, top=THICK, bottom=THICK)
    ws.cell(row=row, column=c).border = Border(left=NONE,  right=NONE, top=THICK, bottom=THICK)
    ws.cell(row=row, column=d).border = Border(left=NONE,  right=THICK, top=THICK, bottom=THICK)
    ws.cell(row=row, column=c).fill   = F_EQ_HDR
    ws.cell(row=row, column=d).fill   = F_EQ_HDR

# ── Load workbook ─────────────────────────────────────────────────────────
print(f"Applying formatting to:\n  {out_file}")
wb = openpyxl.load_workbook(out_file)

for sa in wb.sheetnames:
    print(f"  Formatting sheet: {sa}")
    ws = wb[sa]

    # Column widths
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 10
    for j in range(N_EQ):
        ws.column_dimensions[get_column_letter(col_o(j))].width = 9
        ws.column_dimensions[get_column_letter(col_c(j))].width = 9
        ws.column_dimensions[get_column_letter(col_d(j))].width = 9

    # ── ROW 1: title ──────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 22
    wc(ws, 1, 1, fl=F_TITLE, ft=FT_TITLE, al=align("left","center"))
    for c in range(2, TOTAL_COLS+1):
        ws.cell(row=1, column=c).fill = F_TITLE

    # ── ROW 2: EQ name merged headers ─────────────────────────────────────
    ws.row_dimensions[2].height = 30
    # Fixed cols A, B
    for c in [1, 2]:
        wc(ws, 2, c, fl=F_EQ_HDR, ft=FT_EQ_HDR, al=align())
    # EQ groups
    for j in range(N_EQ):
        o, c, d = col_o(j), col_c(j), col_d(j)
        ws.merge_cells(start_row=2, start_column=o, end_row=2, end_column=d)
        wc(ws, 2, o, fl=F_EQ_HDR, ft=FT_EQ_HDR, al=align(wrap=True))
        apply_eq_merged_border(ws, 2, j)

    # ── ROW 3: sub-column labels ──────────────────────────────────────────
    ws.row_dimensions[3].height = 16
    for c in [1, 2]:
        wc(ws, 3, c, fl=F_EQ_HDR)
    for j in range(N_EQ):
        wc(ws, 3, col_o(j), fl=F_ORIG_HDR, ft=FT_SUB_HDR, al=align(), bd=bd_orig())
        wc(ws, 3, col_c(j), fl=F_CURT_HDR, ft=FT_SUB_HDR, al=align(), bd=bd_curt())
        wc(ws, 3, col_d(j), fl=F_DELT_HDR, ft=FT_SUB_HDR, al=align(), bd=bd_delta())

    # ── ROW 4: unit labels ────────────────────────────────────────────────
    ws.row_dimensions[4].height = 14
    for c in [1, 2]:
        wc(ws, 4, c, fl=F_EQ_HDR)
    for j in range(N_EQ):
        wc(ws, 4, col_o(j), fl=F_ORIG_HDR, ft=FT_SUB_HDR, al=align(), bd=bd_orig())
        wc(ws, 4, col_c(j), fl=F_CURT_HDR, ft=FT_SUB_HDR, al=align(), bd=bd_curt())
        wc(ws, 4, col_d(j), fl=F_DELT_HDR, ft=FT_SUB_HDR, al=align(), bd=bd_delta())

    # ── DATA ROWS (row 5 onward) ───────────────────────────────────────────
    for row_idx, row in enumerate(ws.iter_rows(min_row=5), start=5):
        cell_a = ws.cell(row=row_idx, column=1)
        val_a  = cell_a.value

        # EDP section header row (col A has EDP name, rest merged/blank)
        if val_a and str(val_a).strip() in ['IDR (%)', 'RDR (%)', 'PFA (g)']:
            ws.row_dimensions[row_idx].height = 16
            ws.merge_cells(start_row=row_idx, start_column=1,
                           end_row=row_idx,   end_column=TOTAL_COLS)
            wc(ws, row_idx, 1,
               fl=F_EDP_SEC, ft=FT_EDP_SEC, al=align("left","center"))
            continue

        # Blank separator rows
        if val_a is None and ws.cell(row=row_idx, column=2).value is None:
            continue

        # Data rows — col A blank, col B story/floor label
        ws.row_dimensions[row_idx].height = 14
        wc(ws, row_idx, 1, fl=F_STORY)
        wc(ws, row_idx, 2, fl=F_STORY, ft=FT_STORY, al=align())

        for j in range(N_EQ):
            o, c, d = col_o(j), col_c(j), col_d(j)

            # Original — white
            wc(ws, row_idx, o, fl=F_WHITE, ft=FT_DATA, al=align(),
               fmt='0.000000', bd=bd_orig())

            # Curtailed — white
            wc(ws, row_idx, c, fl=F_WHITE, ft=FT_DATA, al=align(),
               fmt='0.000000', bd=bd_curt())

            # Delta — colour coded
            d_val = ws.cell(row=row_idx, column=d).value
            try:
                d_num = float(d_val)
            except (TypeError, ValueError):
                d_num = None
            wc(ws, row_idx, d, fl=delta_fill(d_num), ft=FT_DATA, al=align(),
               fmt='0.00', bd=bd_delta())

wb.save(out_file)
print("Formatting complete.")
